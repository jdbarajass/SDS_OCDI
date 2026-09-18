"""
Control de Compensatorios Fin de Año — Resolución 2307 de 2026 (modifica
temporalmente la Resolución 2316 de 2023).

Autogestión: cada funcionario registra sus propias horas compensadas (igual
que la Matriz de Seguimiento de abogados, pero aquí aplica a TODOS los
roles, no solo abogados). Admin/jefe pueden ver y editar la de cualquiera,
y son los únicos que pueden ajustar metas por justa causa o configurar el
ciclo (turnos, sábados habilitados, fechas, metas de horas) — todo eso vive
en datos, no en código, para poder reutilizar el módulo en años siguientes.

Reglas de la resolución que modela este módulo:
- 3 turnos de descanso de fin de año, voluntarios, 34h a compensar cada uno.
- Compensación de esas 34h: 1h extra diaria o sábados habilitados, dentro de
  una ventana fija (independiente del turno elegido).
- 24 y 31 de diciembre: jornada corta opcional (7am-2pm), 1.5h fijas de
  compensación por cada fecha trabajada, en una ventana aparte y posterior.
"""
import io
from datetime import date
from fastapi import APIRouter, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from pathlib import Path
from app.template_utils import make_templates

from app.database import get_db
from app.dias_habiles import dias_habiles_diff
from app.auth_utils import (
    tpl, puede_escribir as _pw, registrar_log, historial_registro, ROLES_SUPERUSUARIO,
)

_MOD = "compensatorios"
router = APIRouter(prefix="/compensatorios")
templates = make_templates(str(Path(__file__).parent.parent / "templates"))

TIPOS_REGISTRO = [
    ("HORA_EXTRA_AM", "Hora extra — antes de la jornada"),
    ("HORA_EXTRA_PM", "Hora extra — después de la jornada"),
    ("SABADO",        "Jornada de sábado"),
    ("DIC24",         "24 de diciembre — jornada corta (1.5h fijas)"),
    ("DIC31",         "31 de diciembre — jornada corta (1.5h fijas)"),
]
TIPOS_DIC_FIJOS = {"DIC24": 1.5, "DIC31": 1.5}


# ── Helpers de datos ─────────────────────────────────────────────────────────

def _ciclo_activo(conn) -> dict | None:
    row = conn.execute("SELECT * FROM comp_ciclos WHERE activo=1 ORDER BY id DESC LIMIT 1").fetchone()
    return dict(row) if row else None


def _turnos_ciclo(conn, ciclo_id: int) -> list[dict]:
    rows = conn.execute(
        "SELECT * FROM comp_turnos WHERE ciclo_id=? ORDER BY orden, fecha_inicio", (ciclo_id,)
    ).fetchall()
    return [dict(r) for r in rows]


def _sabados_ciclo(conn, ciclo_id: int) -> list[str]:
    rows = conn.execute(
        "SELECT fecha FROM comp_sabados WHERE ciclo_id=? ORDER BY fecha", (ciclo_id,)
    ).fetchall()
    return [r["fecha"] for r in rows]


def _roster_compensatorios(conn) -> list[str]:
    """Nombres de personal activo con visibilidad en el módulo compensatorios.
    A diferencia de otros módulos (que usan get_personal_oficina() sin filtrar),
    aquí el roster respeta el permiso puede_ver por persona — así admin/jefe
    pueden sacar a alguien del módulo (p.ej. contratistas, a quienes no aplica
    la Resolución 2307, o gente que ya no está en la oficina) desde el propio
    panel de Usuarios y Permisos, sin tocar su acceso al resto del sistema."""
    rows = conn.execute("""
        SELECT u.nombre_completo FROM usuarios u
        LEFT JOIN permisos_modulo p ON p.user_id = u.id AND p.modulo = 'compensatorios'
        WHERE u.activo = 1
          AND (u.rol IN ('admin','jefe') OR COALESCE(p.puede_ver, 1) != 0)
        ORDER BY u.nombre_completo
    """).fetchall()
    return [r[0] for r in rows]


def _nombre_efectivo(user: dict, nombre_qs: str) -> str:
    """Superusuarios pueden operar sobre cualquier funcionario vía ?nombre=;
    todos los demás quedan forzados a su propio nombre sin importar lo que
    venga en el parámetro (mismo patrón de aislamiento que Matriz)."""
    if user.get("rol") in ROLES_SUPERUSUARIO and nombre_qs.strip():
        return nombre_qs.strip().upper()
    return user["nombre_completo"]


def _es_propietario(user: dict | None, nombre_fila: str | None) -> bool:
    if not user:
        return False
    if user.get("rol") in ROLES_SUPERUSUARIO:
        return True
    return bool(nombre_fila) and nombre_fila.strip().upper() == (user.get("nombre_completo") or "").strip().upper()


def _estado_meta(horas_reg: float, meta: float, fecha_limite_str: str | None) -> dict:
    """Semáforo de una meta de horas, reutilizando las mismas clases CSS que
    ya existen para vencimientos de casos (alerta-vencido/proximo/vigente/sin-plazo)."""
    horas_reg = round(horas_reg or 0, 2)
    meta = round(meta or 0, 2)
    if meta <= 0:
        return {"clase": "sin-plazo", "texto": "Sin meta definida", "faltante": 0, "pct": 100}
    pct = min(100, round((horas_reg / meta) * 100))
    faltante = round(meta - horas_reg, 2)
    if faltante <= 0:
        return {"clase": "vigente", "texto": f"Completo · {horas_reg:g}/{meta:g} h", "faltante": 0, "pct": 100}
    if not fecha_limite_str:
        return {"clase": "proximo", "texto": f"Faltan {faltante:g} h de {meta:g} h", "faltante": faltante, "pct": pct}
    try:
        fecha_limite = date.fromisoformat(str(fecha_limite_str)[:10])
    except ValueError:
        return {"clase": "proximo", "texto": f"Faltan {faltante:g} h de {meta:g} h", "faltante": faltante, "pct": pct}
    hoy = date.today()
    if fecha_limite < hoy:
        return {"clase": "vencido", "texto": f"Venció el plazo — faltan {faltante:g} h", "faltante": faltante, "pct": pct}
    dias_restantes = dias_habiles_diff(hoy, fecha_limite)
    if dias_restantes < faltante:
        return {"clase": "vencido",
                "texto": f"En riesgo — faltan {faltante:g} h en {dias_restantes} días hábiles",
                "faltante": faltante, "pct": pct}
    return {"clase": "proximo",
            "texto": f"Faltan {faltante:g} h de {meta:g} h · {dias_restantes} días hábiles disponibles",
            "faltante": faltante, "pct": pct}


def _resumen_funcionario(conn, ciclo: dict, nombre: str) -> dict:
    fila = conn.execute(
        "SELECT * FROM comp_funcionarios WHERE ciclo_id=? AND nombre_completo=?",
        (ciclo["id"], nombre),
    ).fetchone()
    fila = dict(fila) if fila else {
        "id": None, "turno_id": None, "participa": 1, "tiene_compensatorios_previos": 0,
        "ajuste_horas_turno": 0, "ajuste_horas_dic": 0, "motivo_ajuste": None, "observaciones": None,
    }

    turno = None
    if fila.get("turno_id"):
        t = conn.execute("SELECT * FROM comp_turnos WHERE id=?", (fila["turno_id"],)).fetchone()
        turno = dict(t) if t else None

    horas_turno = conn.execute(
        "SELECT COALESCE(SUM(horas),0) FROM comp_registro "
        "WHERE ciclo_id=? AND nombre_completo=? AND aplica_a='TURNO' AND eliminado_en IS NULL",
        (ciclo["id"], nombre),
    ).fetchone()[0]
    horas_dic = conn.execute(
        "SELECT COALESCE(SUM(horas),0) FROM comp_registro "
        "WHERE ciclo_id=? AND nombre_completo=? AND aplica_a='DIC' AND eliminado_en IS NULL",
        (ciclo["id"], nombre),
    ).fetchone()[0]

    meta_turno = max(0.0, (ciclo.get("meta_horas_turno") or 0) + (fila.get("ajuste_horas_turno") or 0))
    meta_dic = max(0.0, (ciclo.get("meta_horas_dic") or 0) + (fila.get("ajuste_horas_dic") or 0))

    estado_turno = _estado_meta(horas_turno, meta_turno, ciclo.get("ventana_fin"))
    estado_dic = _estado_meta(horas_dic, meta_dic, ciclo.get("ventana_dic_fin"))

    hoy = date.today()
    descansa_hoy = False
    if turno and fila.get("participa"):
        try:
            fi = date.fromisoformat(turno["fecha_inicio"])
            ff = date.fromisoformat(turno["fecha_fin"])
            descansa_hoy = fi <= hoy <= ff
        except (ValueError, TypeError):
            pass

    resumen = dict(fila)
    resumen.update({
        "nombre_completo": nombre,
        "turno": turno,
        "horas_turno": horas_turno, "meta_turno": meta_turno, "estado_turno": estado_turno,
        "horas_dic": horas_dic, "meta_dic": meta_dic, "estado_dic": estado_dic,
        "descansa_hoy": descansa_hoy,
    })
    return resumen


def _require_super(request: Request):
    user = getattr(request.state, "user", None)
    if not user or user.get("rol") not in ROLES_SUPERUSUARIO:
        return None
    return user


# ── Dashboard ────────────────────────────────────────────────────────────────

@router.get("/", response_class=HTMLResponse)
async def dashboard(request: Request, msg: str = ""):
    user = request.state.user
    conn = get_db()
    ciclo = _ciclo_activo(conn)
    if not ciclo:
        conn.close()
        return templates.TemplateResponse("compensatorios_dashboard.html", tpl(request, _MOD,
            ciclo=None, turnos=[], resumenes=[], mi_resumen=None, por_turno={},
            active="compensatorios_dashboard", msg=msg,
        ))

    turnos = _turnos_ciclo(conn, ciclo["id"])
    nombres = _roster_compensatorios(conn)
    resumenes = [_resumen_funcionario(conn, ciclo, n) for n in nombres]
    conn.close()

    mi_resumen = next((r for r in resumenes if r["nombre_completo"] == user["nombre_completo"]), None)

    por_turno: dict[int, list] = {t["id"]: [] for t in turnos}
    for r in resumenes:
        if r.get("turno") and r.get("participa"):
            por_turno.setdefault(r["turno"]["id"], []).append(r["nombre_completo"])

    return templates.TemplateResponse("compensatorios_dashboard.html", tpl(request, _MOD,
        ciclo=ciclo, turnos=turnos, resumenes=resumenes, mi_resumen=mi_resumen, por_turno=por_turno,
        active="compensatorios_dashboard", msg=msg,
    ))


# ── Mis horas (autogestión) ───────────────────────────────────────────────────

@router.get("/mis-horas", response_class=HTMLResponse)
async def mis_horas(request: Request, nombre: str = "", msg: str = ""):
    user = request.state.user
    conn = get_db()
    ciclo = _ciclo_activo(conn)
    if not ciclo:
        conn.close()
        return RedirectResponse("/compensatorios/?msg=sin_ciclo", status_code=303)

    objetivo = _nombre_efectivo(user, nombre)
    turnos = _turnos_ciclo(conn, ciclo["id"])
    sabados = _sabados_ciclo(conn, ciclo["id"])
    resumen = _resumen_funcionario(conn, ciclo, objetivo)
    registros = conn.execute(
        "SELECT * FROM comp_registro WHERE ciclo_id=? AND nombre_completo=? AND eliminado_en IS NULL "
        "ORDER BY fecha DESC, id DESC",
        (ciclo["id"], objetivo),
    ).fetchall()
    conn.close()

    return templates.TemplateResponse("compensatorios_mis_horas.html", tpl(request, _MOD,
        ciclo=ciclo, turnos=turnos, sabados=sabados, resumen=resumen,
        registros=[dict(r) for r in registros], objetivo=objetivo, tipos=TIPOS_REGISTRO,
        tipos_dict=dict(TIPOS_REGISTRO), hoy=date.today().isoformat(),
        es_superuser_viendo_otro=(user.get("rol") in ROLES_SUPERUSUARIO and objetivo != user["nombre_completo"]),
        active="compensatorios_mis_horas", msg=msg,
    ))


@router.post("/mis-horas/turno")
async def guardar_turno(
    request: Request,
    nombre: str = Form(""),
    turno_id: str = Form(""),
    participa: str = Form(""),
    tiene_compensatorios_previos: str = Form(""),
    observaciones: str = Form(""),
    ajuste_horas_turno: str = Form("0"),
    ajuste_horas_dic: str = Form("0"),
    motivo_ajuste: str = Form(""),
):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    ciclo = _ciclo_activo(conn)
    if not ciclo:
        conn.close()
        return RedirectResponse("/compensatorios/?msg=sin_ciclo", status_code=303)

    objetivo = _nombre_efectivo(user, nombre)
    tid = int(turno_id) if turno_id.strip().isdigit() else None
    part = 1 if participa in ("on", "1") else 0
    prev = 1 if tiene_compensatorios_previos in ("on", "1") else 0

    es_super = user.get("rol") in ROLES_SUPERUSUARIO
    existente = conn.execute(
        "SELECT ajuste_horas_turno, ajuste_horas_dic, motivo_ajuste FROM comp_funcionarios "
        "WHERE ciclo_id=? AND nombre_completo=?",
        (ciclo["id"], objetivo),
    ).fetchone()

    if es_super:
        try:
            ajuste_t = float(ajuste_horas_turno.replace(",", "."))
        except ValueError:
            ajuste_t = existente["ajuste_horas_turno"] if existente else 0
        try:
            ajuste_d = float(ajuste_horas_dic.replace(",", "."))
        except ValueError:
            ajuste_d = existente["ajuste_horas_dic"] if existente else 0
        motivo = motivo_ajuste.strip() or None
    else:
        # Los campos de ajuste solo los puede tocar admin/jefe — si un no-superusuario
        # llega a spoofear el POST, se ignoran y se conservan los valores existentes.
        ajuste_t = existente["ajuste_horas_turno"] if existente else 0
        ajuste_d = existente["ajuste_horas_dic"] if existente else 0
        motivo = existente["motivo_ajuste"] if existente else None

    conn.execute("""
        INSERT INTO comp_funcionarios
            (ciclo_id, nombre_completo, turno_id, participa, tiene_compensatorios_previos,
             ajuste_horas_turno, ajuste_horas_dic, motivo_ajuste, observaciones)
        VALUES (?,?,?,?,?,?,?,?,?)
        ON CONFLICT(ciclo_id, nombre_completo) DO UPDATE SET
            turno_id=excluded.turno_id, participa=excluded.participa,
            tiene_compensatorios_previos=excluded.tiene_compensatorios_previos,
            ajuste_horas_turno=excluded.ajuste_horas_turno,
            ajuste_horas_dic=excluded.ajuste_horas_dic,
            motivo_ajuste=excluded.motivo_ajuste,
            observaciones=excluded.observaciones,
            updated_at=datetime('now','localtime')
    """, (ciclo["id"], objetivo, tid, part, prev, ajuste_t, ajuste_d, motivo, observaciones.strip() or None))
    conn.commit()
    conn.close()
    registrar_log(user, "editar", _MOD, f"{objetivo} — turno/ajustes de compensatorios actualizados")
    return RedirectResponse(f"/compensatorios/mis-horas?nombre={objetivo}&msg=turno_actualizado", status_code=303)


# ── Registro de horas ─────────────────────────────────────────────────────────

@router.post("/registro/nuevo")
async def registro_nuevo(
    request: Request,
    nombre: str = Form(""),
    fecha: str = Form(""),
    tipo: str = Form(""),
    horas: str = Form("0"),
    actividad: str = Form(""),
    observaciones: str = Form(""),
):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    ciclo = _ciclo_activo(conn)
    if not ciclo:
        conn.close()
        return RedirectResponse("/compensatorios/?msg=sin_ciclo", status_code=303)

    objetivo = _nombre_efectivo(user, nombre)
    tipo = tipo.strip().upper()
    if tipo not in dict(TIPOS_REGISTRO) or not fecha:
        conn.close()
        return RedirectResponse(f"/compensatorios/mis-horas?nombre={objetivo}&msg=error_datos", status_code=303)

    aplica_a = "DIC" if tipo in TIPOS_DIC_FIJOS else "TURNO"

    if tipo in TIPOS_DIC_FIJOS:
        dup = conn.execute(
            "SELECT 1 FROM comp_registro WHERE ciclo_id=? AND nombre_completo=? AND tipo=? AND eliminado_en IS NULL",
            (ciclo["id"], objetivo, tipo),
        ).fetchone()
        if dup:
            conn.close()
            return RedirectResponse(f"/compensatorios/mis-horas?nombre={objetivo}&msg=duplicado", status_code=303)
        horas_val = TIPOS_DIC_FIJOS[tipo]
    else:
        try:
            horas_val = round(float(horas.replace(",", ".")), 2)
        except ValueError:
            horas_val = 0
        if horas_val <= 0:
            conn.close()
            return RedirectResponse(f"/compensatorios/mis-horas?nombre={objetivo}&msg=error_datos", status_code=303)

    cur = conn.execute(
        """INSERT INTO comp_registro
           (ciclo_id, nombre_completo, fecha, tipo, horas, aplica_a, actividad, observaciones, created_by)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (ciclo["id"], objetivo, fecha, tipo, horas_val, aplica_a,
         actividad.strip() or None, observaciones.strip() or None, user.get("nombre_completo")),
    )
    new_id = cur.lastrowid
    conn.commit()
    conn.close()
    registrar_log(user, "crear", _MOD, f"{objetivo} — {tipo} — {fecha} — {horas_val:g}h",
                  request.client.host if request.client else None, registro_id=new_id)
    return RedirectResponse(f"/compensatorios/mis-horas?nombre={objetivo}&msg=creado", status_code=303)


@router.get("/registro/{reg_id}/editar", response_class=HTMLResponse)
async def registro_editar_form(request: Request, reg_id: int):
    user = request.state.user
    conn = get_db()
    reg = conn.execute("SELECT * FROM comp_registro WHERE id=?", (reg_id,)).fetchone()
    if not reg or not _es_propietario(user, reg["nombre_completo"]):
        conn.close()
        return RedirectResponse("/compensatorios/?msg=no_encontrado")
    if reg["eliminado_en"]:
        conn.close()
        return RedirectResponse("/compensatorios/papelera?msg=error_en_papelera", status_code=303)
    if not _pw(user, _MOD):
        conn.close()
        return RedirectResponse(f"/compensatorios/mis-horas?nombre={reg['nombre_completo']}&msg=sin_permiso", status_code=303)
    historial = historial_registro(conn, _MOD, reg_id)
    conn.close()
    return templates.TemplateResponse("compensatorios_registro_editar.html", tpl(request, _MOD,
        reg=dict(reg), tipos=TIPOS_REGISTRO, tipos_dict=dict(TIPOS_REGISTRO),
        historial=historial, active="compensatorios_mis_horas",
    ))


@router.post("/registro/{reg_id}/editar")
async def registro_editar_post(
    request: Request,
    reg_id: int,
    fecha: str = Form(""),
    tipo: str = Form(""),
    horas: str = Form("0"),
    actividad: str = Form(""),
    observaciones: str = Form(""),
):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse(f"/compensatorios/registro/{reg_id}/editar?msg=sin_permiso", status_code=303)
    conn = get_db()
    actual = conn.execute(
        "SELECT nombre_completo, tipo FROM comp_registro WHERE id=? AND eliminado_en IS NULL", (reg_id,)
    ).fetchone()
    if not actual or not _es_propietario(user, actual["nombre_completo"]):
        conn.close()
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)

    tipo = tipo.strip().upper()
    if tipo not in dict(TIPOS_REGISTRO) or not fecha:
        conn.close()
        return RedirectResponse(f"/compensatorios/registro/{reg_id}/editar?msg=error_datos", status_code=303)
    aplica_a = "DIC" if tipo in TIPOS_DIC_FIJOS else "TURNO"

    if tipo in TIPOS_DIC_FIJOS:
        horas_val = TIPOS_DIC_FIJOS[tipo]
    else:
        try:
            horas_val = round(float(horas.replace(",", ".")), 2)
        except ValueError:
            horas_val = 0
        if horas_val <= 0:
            conn.close()
            return RedirectResponse(f"/compensatorios/registro/{reg_id}/editar?msg=error_datos", status_code=303)

    conn.execute(
        """UPDATE comp_registro SET fecha=?, tipo=?, horas=?, aplica_a=?, actividad=?, observaciones=?,
               updated_at=datetime('now','localtime')
           WHERE id=? AND eliminado_en IS NULL""",
        (fecha, tipo, horas_val, aplica_a, actividad.strip() or None, observaciones.strip() or None, reg_id),
    )
    conn.commit()
    conn.close()
    registrar_log(user, "editar", _MOD, f"Registro #{reg_id} actualizado", registro_id=reg_id)
    return RedirectResponse(f"/compensatorios/mis-horas?nombre={actual['nombre_completo']}&msg=actualizado", status_code=303)


@router.post("/registro/{reg_id}/eliminar")
async def registro_eliminar(request: Request, reg_id: int):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    actual = conn.execute("SELECT nombre_completo FROM comp_registro WHERE id=?", (reg_id,)).fetchone()
    if not actual or not _es_propietario(user, actual["nombre_completo"]):
        conn.close()
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn.execute(
        "UPDATE comp_registro SET eliminado_en=datetime('now','localtime'), eliminado_por=? WHERE id=?",
        (user.get("nombre_completo"), reg_id),
    )
    conn.commit()
    conn.close()
    registrar_log(user, "eliminar", _MOD, f"Registro #{reg_id}",
                  request.client.host if request.client else None, registro_id=reg_id)
    return RedirectResponse(f"/compensatorios/mis-horas?nombre={actual['nombre_completo']}&msg=eliminado", status_code=303)


# ── Papelera de reciclaje ──────────────────────────────────────────────────────

@router.get("/papelera", response_class=HTMLResponse)
async def papelera(request: Request, msg: str = ""):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    where = ["eliminado_en IS NOT NULL"]
    params: list = []
    if user.get("rol") not in ROLES_SUPERUSUARIO:
        where.append("nombre_completo = ?")
        params.append(user["nombre_completo"])
    rows = conn.execute(
        f"SELECT * FROM comp_registro WHERE {' AND '.join(where)} ORDER BY eliminado_en DESC", params,
    ).fetchall()
    conn.close()
    registros = [{
        "id": r["id"],
        "titulo": f"{dict(TIPOS_REGISTRO).get(r['tipo'], r['tipo'])} — {r['fecha']} ({r['horas']:g}h)",
        "subtitulo": r["nombre_completo"] or "",
        "eliminado_en": r["eliminado_en"],
        "eliminado_por": r["eliminado_por"],
    } for r in rows]
    return templates.TemplateResponse("papelera.html", tpl(request, _MOD,
        modulo_nombre="Compensatorios Fin de Año", prefix="/compensatorios/registro",
        base_template="base_compensatorios.html", registros=registros,
        volver_url="/compensatorios/mis-horas", active="papelera", msg=msg,
    ))


@router.post("/registro/{reg_id}/restaurar")
async def restaurar(request: Request, reg_id: int):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/compensatorios/papelera?msg=sin_permiso", status_code=303)
    conn = get_db()
    row = conn.execute("SELECT nombre_completo FROM comp_registro WHERE id=?", (reg_id,)).fetchone()
    if not row or not _es_propietario(user, row["nombre_completo"]):
        conn.close()
        return RedirectResponse("/compensatorios/papelera?msg=sin_permiso", status_code=303)
    conn.execute("UPDATE comp_registro SET eliminado_en=NULL, eliminado_por=NULL WHERE id=?", (reg_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "restaurar", _MOD, f"Registro #{reg_id}", registro_id=reg_id)
    return RedirectResponse("/compensatorios/papelera?msg=restaurado", status_code=303)


@router.post("/registro/{reg_id}/purgar")
async def purgar(request: Request, reg_id: int):
    user = request.state.user
    if not user or user.get("rol") not in ROLES_SUPERUSUARIO:
        return RedirectResponse("/compensatorios/papelera?msg=sin_permiso", status_code=303)
    conn = get_db()
    conn.execute("DELETE FROM comp_registro WHERE id=? AND eliminado_en IS NOT NULL", (reg_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "purgar", _MOD, f"Registro #{reg_id} — eliminado definitivamente", registro_id=reg_id)
    return RedirectResponse("/compensatorios/papelera?msg=purgado", status_code=303)


# ── Exportar ───────────────────────────────────────────────────────────────────

@router.get("/exportar")
async def exportar(request: Request):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return RedirectResponse("/compensatorios/?msg=error_openpyxl")

    conn = get_db()
    ciclo = _ciclo_activo(conn)
    if not ciclo:
        conn.close()
        return RedirectResponse("/compensatorios/?msg=sin_ciclo")
    nombres = _roster_compensatorios(conn)
    resumenes = [_resumen_funcionario(conn, ciclo, n) for n in nombres]
    roles = {r["nombre_completo"]: r["rol"] for r in conn.execute("SELECT nombre_completo, rol FROM usuarios").fetchall()}
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COMPENSATORIOS"

    headers = [
        "FUNCIONARIO", "ROL", "TURNO ELEGIDO", "FECHAS DEL TURNO", "PARTICIPA",
        "HORAS TURNO", "META TURNO", "% TURNO", "ESTADO TURNO",
        "HORAS DIC 24/31", "META DIC", "% DIC", "ESTADO DIC",
        "COMPENSATORIOS PREVIOS", "OBSERVACIONES",
    ]
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor="0D3060")

    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = center
        c.border = border
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    _na = lambda v: v if (v is not None and str(v).strip() != "") else "N/A"

    for r in resumenes:
        turno = r.get("turno")
        row_data = [
            r["nombre_completo"], _na(roles.get(r["nombre_completo"])),
            _na(turno["nombre"] if turno else None),
            _na(f"{turno['fecha_inicio']} a {turno['fecha_fin']}" if turno else None),
            "SI" if r.get("participa") else "NO",
            r["horas_turno"], r["meta_turno"], r["estado_turno"]["pct"], r["estado_turno"]["texto"],
            r["horas_dic"], r["meta_dic"], r["estado_dic"]["pct"], r["estado_dic"]["texto"],
            "SI" if r.get("tiene_compensatorios_previos") else "NO",
            _na(r.get("observaciones")),
        ]
        ws.append(row_data)
        ri = ws.max_row
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.border = border
            cell.alignment = wrap
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    anchos = [28, 13, 14, 22, 10, 12, 11, 9, 30, 14, 10, 8, 30, 14, 30]
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    hoy = date.today().strftime("%Y%m%d")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Compensatorios_FinDeAno_{hoy}.xlsx"},
    )


# ── Configuración del ciclo (solo admin/jefe) ─────────────────────────────────

@router.get("/ciclo", response_class=HTMLResponse)
async def ciclo_config(request: Request, msg: str = ""):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    ciclos = conn.execute("SELECT * FROM comp_ciclos ORDER BY id DESC").fetchall()
    ciclo = _ciclo_activo(conn)
    turnos = _turnos_ciclo(conn, ciclo["id"]) if ciclo else []
    sabados = conn.execute(
        "SELECT * FROM comp_sabados WHERE ciclo_id=? ORDER BY fecha", (ciclo["id"],)
    ).fetchall() if ciclo else []
    conn.close()
    return templates.TemplateResponse("compensatorios_ciclo.html", tpl(request, _MOD,
        ciclos=[dict(c) for c in ciclos], ciclo=ciclo, turnos=turnos,
        sabados=[dict(s) for s in sabados], active="compensatorios_ciclo", msg=msg,
    ))


@router.post("/ciclo/nuevo")
async def ciclo_nuevo(request: Request, nombre: str = Form(...), resolucion: str = Form("")):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    if not nombre.strip():
        return RedirectResponse("/compensatorios/ciclo?msg=error_datos", status_code=303)
    conn = get_db()
    conn.execute("UPDATE comp_ciclos SET activo=0")
    conn.execute(
        "INSERT INTO comp_ciclos (nombre, resolucion, activo, meta_horas_turno, meta_horas_dic, created_by) "
        "VALUES (?,?,1,34,3,?)",
        (nombre.strip(), resolucion.strip() or None, user.get("nombre_completo")),
    )
    conn.commit()
    conn.close()
    registrar_log(user, "crear", _MOD, f"Nuevo ciclo de compensatorios: '{nombre}'")
    return RedirectResponse("/compensatorios/ciclo?msg=ciclo_creado", status_code=303)


@router.post("/ciclo/{ciclo_id}/activar")
async def ciclo_activar(request: Request, ciclo_id: int):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    conn.execute("UPDATE comp_ciclos SET activo=0")
    conn.execute("UPDATE comp_ciclos SET activo=1 WHERE id=?", (ciclo_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "editar", _MOD, f"Ciclo #{ciclo_id} activado")
    return RedirectResponse("/compensatorios/ciclo?msg=ciclo_activado", status_code=303)


@router.post("/ciclo/editar")
async def ciclo_editar(
    request: Request,
    ciclo_id: int = Form(...),
    nombre: str = Form(""),
    resolucion: str = Form(""),
    meta_horas_turno: str = Form("34"),
    meta_horas_dic: str = Form("3"),
    ventana_ini: str = Form(""),
    ventana_fin: str = Form(""),
    ventana_dic_ini: str = Form(""),
    ventana_dic_fin: str = Form(""),
    fecha_tope_reporte: str = Form(""),
):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    try:
        mht = float(meta_horas_turno.replace(",", "."))
        mhd = float(meta_horas_dic.replace(",", "."))
    except ValueError:
        return RedirectResponse("/compensatorios/ciclo?msg=error_datos", status_code=303)
    conn = get_db()
    conn.execute("""
        UPDATE comp_ciclos SET nombre=?, resolucion=?, meta_horas_turno=?, meta_horas_dic=?,
            ventana_ini=?, ventana_fin=?, ventana_dic_ini=?, ventana_dic_fin=?, fecha_tope_reporte=?
        WHERE id=?
    """, (nombre.strip() or "Ciclo", resolucion.strip() or None, mht, mhd,
          ventana_ini or None, ventana_fin or None, ventana_dic_ini or None, ventana_dic_fin or None,
          fecha_tope_reporte or None, ciclo_id))
    conn.commit()
    conn.close()
    registrar_log(user, "editar", _MOD, f"Ciclo #{ciclo_id} — parámetros actualizados")
    return RedirectResponse("/compensatorios/ciclo?msg=actualizado", status_code=303)


@router.post("/ciclo/turno/nuevo")
async def turno_nuevo(
    request: Request, ciclo_id: int = Form(...), nombre: str = Form(""),
    fecha_inicio: str = Form(""), fecha_fin: str = Form(""), orden: int = Form(0),
):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    if not nombre.strip() or not fecha_inicio or not fecha_fin:
        return RedirectResponse("/compensatorios/ciclo?msg=error_datos", status_code=303)
    conn = get_db()
    conn.execute(
        "INSERT INTO comp_turnos (ciclo_id, nombre, fecha_inicio, fecha_fin, orden) VALUES (?,?,?,?,?)",
        (ciclo_id, nombre.strip(), fecha_inicio, fecha_fin, orden),
    )
    conn.commit()
    conn.close()
    registrar_log(user, "crear", _MOD, f"Turno '{nombre}' agregado al ciclo #{ciclo_id}")
    return RedirectResponse("/compensatorios/ciclo?msg=turno_creado", status_code=303)


@router.post("/ciclo/turno/{turno_id}/eliminar")
async def turno_eliminar(request: Request, turno_id: int):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    conn.execute("UPDATE comp_funcionarios SET turno_id=NULL WHERE turno_id=?", (turno_id,))
    conn.execute("DELETE FROM comp_turnos WHERE id=?", (turno_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "eliminar", _MOD, f"Turno #{turno_id} eliminado")
    return RedirectResponse("/compensatorios/ciclo?msg=turno_eliminado", status_code=303)


@router.post("/ciclo/sabado/nuevo")
async def sabado_nuevo(request: Request, ciclo_id: int = Form(...), fecha: str = Form("")):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    if not fecha:
        return RedirectResponse("/compensatorios/ciclo?msg=error_datos", status_code=303)
    conn = get_db()
    conn.execute("INSERT INTO comp_sabados (ciclo_id, fecha) VALUES (?,?)", (ciclo_id, fecha))
    conn.commit()
    conn.close()
    registrar_log(user, "crear", _MOD, f"Sábado habilitado {fecha} agregado al ciclo #{ciclo_id}")
    return RedirectResponse("/compensatorios/ciclo?msg=sabado_creado", status_code=303)


@router.post("/ciclo/sabado/{sabado_id}/eliminar")
async def sabado_eliminar(request: Request, sabado_id: int):
    user = _require_super(request)
    if not user:
        return RedirectResponse("/compensatorios/?msg=sin_permiso", status_code=303)
    conn = get_db()
    conn.execute("DELETE FROM comp_sabados WHERE id=?", (sabado_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "eliminar", _MOD, f"Sábado #{sabado_id} eliminado")
    return RedirectResponse("/compensatorios/ciclo?msg=sabado_eliminado", status_code=303)
