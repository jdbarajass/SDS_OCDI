"""
Matriz de Seguimiento — bitácora manual de cada abogado sobre sus trámites en
BPM (plataforma AgilSalud, donde se radican y suben los expedientes). No es
un espejo automático de otros módulos: el abogado la llena a mano como
recordatorio de en qué etapa va cada trámite y qué le falta por hacer.

Cada abogado solo ve y edita sus propias filas (el filtro por 'abogado' se
fuerza server-side a su propio nombre); admin/jefe supervisan todas.
"""
import io
from datetime import date, datetime
from fastapi import APIRouter, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from pathlib import Path
from app.template_utils import make_templates

from app.database import get_db, calcular_alerta
from app.auth_utils import (
    tpl, puede_escribir as _pw, registrar_log, historial_registro, ROLES_SUPERUSUARIO,
)

_MOD = "matriz"

router = APIRouter(prefix="/matriz")
templates = make_templates(str(Path(__file__).parent.parent / "templates"))

POR_PAGINA = 30

TIPOS_TRAMITE = [
    "INDAGACIÓN PREVIA",
    "INVESTIGACIÓN DISCIPLINARIA",
    "SDQS",
    "CORRESPONDENCIA / DERECHO DE PETICIÓN",
    "TUTELA",
    "EXPEDIENTE DIGITAL",
    "OTRO",
]

ETAPAS_BPM = [
    "RADICADO / SIN ASIGNAR",
    "EN ESTUDIO",
    "EN TRÁMITE",
    "PENDIENTE DOCUMENTOS",
    "PENDIENTE DE FIRMA",
    "PENDIENTE DE NOTIFICACIÓN",
    "EN REVISIÓN JURÍDICA",
    "RESPONDIDO",
    "CERRADO",
    "ARCHIVADO",
]

ESTADOS = ["EN TRÁMITE", "AL DÍA", "PENDIENTE", "VENCIDO", "ARCHIVADO"]


def _abogados_lista(conn) -> list[str]:
    return [r[0] for r in conn.execute(
        "SELECT nombre_completo FROM usuarios WHERE rol='abogado' AND activo=1 ORDER BY nombre_completo"
    ).fetchall()]


def _es_propietario(user: dict | None, abogado_fila: str | None) -> bool:
    if not user:
        return False
    if user.get("rol") in ROLES_SUPERUSUARIO:
        return True
    if user.get("rol") != "abogado":
        return False
    return bool(abogado_fila) and abogado_fila.strip().upper() == (user.get("nombre_completo") or "").strip().upper()


def _enriquecer(row: dict) -> dict:
    row["alerta"] = calcular_alerta(row.get("fecha_limite"))
    return row


# ── Lista ──────────────────────────────────────────────────────────────────────

@router.get("/", response_class=HTMLResponse)
async def matriz_lista(
    request: Request,
    q: str = "",
    abogado: str = "",
    tipo_tramite: str = "",
    etapa_bpm: str = "",
    estado: str = "",
    page: int = 1,
    msg: str = "",
):
    user = request.state.user
    conn = get_db()

    es_abogado = bool(user) and user.get("rol") == "abogado"
    if es_abogado:
        abogado = user["nombre_completo"]

    where, params = ["eliminado_en IS NULL"], []
    if abogado:
        where.append("abogado = ?")
        params.append(abogado)
    if q:
        where.append("(n_expediente LIKE ? OR n_bpm LIKE ? OR asunto LIKE ?)")
        params += [f"%{q}%"] * 3
    if tipo_tramite:
        where.append("tipo_tramite = ?")
        params.append(tipo_tramite)
    if etapa_bpm:
        where.append("etapa_bpm = ?")
        params.append(etapa_bpm)
    if estado:
        where.append("estado = ?")
        params.append(estado)

    cond = "WHERE " + " AND ".join(where)
    total = conn.execute(f"SELECT COUNT(*) FROM matriz_seguimiento {cond}", params).fetchone()[0]
    offset = (page - 1) * POR_PAGINA
    rows = conn.execute(
        f"""SELECT * FROM matriz_seguimiento {cond}
            ORDER BY (fecha_limite IS NULL), fecha_limite ASC, id DESC
            LIMIT ? OFFSET ?""",
        params + [POR_PAGINA, offset],
    ).fetchall()

    abogados = _abogados_lista(conn)
    rows_dict = [_enriquecer(dict(r)) for r in rows]
    conn.close()

    total_pages = max(1, (total + POR_PAGINA - 1) // POR_PAGINA)
    return templates.TemplateResponse("matriz_lista.html", tpl(request, _MOD,
        rows=rows_dict, total=total, page=page, total_pages=total_pages,
        q=q, abogado=abogado, tipo_tramite=tipo_tramite, etapa_bpm=etapa_bpm, estado=estado,
        abogados=abogados, tipos=TIPOS_TRAMITE, etapas=ETAPAS_BPM, estados=ESTADOS,
        es_abogado=es_abogado, msg=msg, active="matriz_lista",
    ))


# ── Nuevo ──────────────────────────────────────────────────────────────────────

@router.get("/nuevo", response_class=HTMLResponse)
async def matriz_nuevo_form(request: Request):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/matriz/?msg=sin_permiso", status_code=303)
    conn = get_db()
    abogados = _abogados_lista(conn)
    conn.close()
    es_abogado = user.get("rol") == "abogado"
    return templates.TemplateResponse("matriz_form.html", tpl(request, _MOD,
        reg=None, abogados=abogados, tipos=TIPOS_TRAMITE, etapas=ETAPAS_BPM, estados=ESTADOS,
        es_abogado=es_abogado, active="matriz_nuevo",
    ))


@router.post("/nuevo")
async def matriz_nuevo_post(
    request: Request,
    abogado: str = Form(""),
    n_expediente: str = Form(""),
    n_bpm: str = Form(""),
    tipo_tramite: str = Form(""),
    etapa_bpm: str = Form(""),
    asunto: str = Form(""),
    ultima_actuacion: str = Form(""),
    fecha_ultima_actuacion: str = Form(""),
    proxima_actuacion: str = Form(""),
    fecha_limite: str = Form(""),
    estado: str = Form(""),
    observaciones: str = Form(""),
):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/matriz/?msg=sin_permiso", status_code=303)

    # El abogado solo puede crear filas de su propia matriz, sin importar
    # qué valor venga en el formulario.
    if user.get("rol") == "abogado":
        abogado = user["nombre_completo"]
    abogado = abogado.strip()
    if not abogado:
        return RedirectResponse("/matriz/nuevo?msg=error_datos", status_code=303)

    conn = get_db()
    cur = conn.execute(
        """INSERT INTO matriz_seguimiento
           (abogado, n_expediente, n_bpm, tipo_tramite, etapa_bpm, asunto,
            ultima_actuacion, fecha_ultima_actuacion, proxima_actuacion, fecha_limite,
            estado, observaciones, created_by)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        [
            abogado,
            n_expediente.strip() or None,
            n_bpm.strip() or None,
            tipo_tramite.strip() or None,
            etapa_bpm.strip() or None,
            asunto.strip() or None,
            ultima_actuacion.strip() or None,
            fecha_ultima_actuacion or None,
            proxima_actuacion.strip() or None,
            fecha_limite or None,
            estado.strip() or "EN TRÁMITE",
            observaciones.strip() or None,
            user.get("nombre_completo"),
        ],
    )
    new_id = cur.lastrowid
    conn.commit()
    conn.close()
    registrar_log(user, "crear", _MOD, f"{abogado} — BPM {n_bpm or 's/n'} — {asunto or n_expediente or ''}",
                  request.client.host if request.client else None, registro_id=new_id)
    return RedirectResponse("/matriz/?msg=creado", status_code=303)


# ── Exportar ───────────────────────────────────────────────────────────────────

@router.get("/exportar")
async def matriz_exportar(
    request: Request,
    q: str = "",
    abogado: str = "",
    tipo_tramite: str = "",
    etapa_bpm: str = "",
    estado: str = "",
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return RedirectResponse("/matriz/?msg=error_openpyxl")

    user = request.state.user
    if user.get("rol") == "abogado":
        abogado = user["nombre_completo"]

    where, params = ["eliminado_en IS NULL"], []
    if abogado:
        where.append("abogado = ?")
        params.append(abogado)
    if q:
        where.append("(n_expediente LIKE ? OR n_bpm LIKE ? OR asunto LIKE ?)")
        params += [f"%{q}%"] * 3
    if tipo_tramite:
        where.append("tipo_tramite = ?")
        params.append(tipo_tramite)
    if etapa_bpm:
        where.append("etapa_bpm = ?")
        params.append(etapa_bpm)
    if estado:
        where.append("estado = ?")
        params.append(estado)
    cond = "WHERE " + " AND ".join(where)

    conn = get_db()
    rows = conn.execute(
        f"SELECT * FROM matriz_seguimiento {cond} ORDER BY abogado, (fecha_limite IS NULL), fecha_limite ASC",
        params,
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MATRIZ SEGUIMIENTO"

    headers = [
        "ABOGADO", "N° EXPEDIENTE", "N° BPM", "TIPO DE TRÁMITE", "ETAPA BPM", "ASUNTO",
        "ÚLTIMA ACTUACIÓN", "FECHA ÚLTIMA ACTUACIÓN", "PRÓXIMA ACTUACIÓN", "FECHA LÍMITE",
        "ESTADO", "OBSERVACIONES",
    ]
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor="0D3060")

    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = center
        c.border = border
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    _na = lambda v: v if (v is not None and str(v).strip() != "") else "N/A"

    for r in rows:
        d = dict(r)
        row_data = [
            _na(d.get("abogado")), _na(d.get("n_expediente")), _na(d.get("n_bpm")),
            _na(d.get("tipo_tramite")), _na(d.get("etapa_bpm")), _na(d.get("asunto")),
            _na(d.get("ultima_actuacion")), _na(d.get("fecha_ultima_actuacion")),
            _na(d.get("proxima_actuacion")), _na(d.get("fecha_limite")),
            _na(d.get("estado")), _na(d.get("observaciones")),
        ]
        ws.append(row_data)
        ri = ws.max_row
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.border = border
            cell.alignment = wrap
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    anchos = [26, 14, 14, 22, 20, 32, 28, 16, 28, 14, 14, 30]
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    hoy = date.today().strftime("%Y%m%d")
    sufijo = f"_{abogado.split()[0]}" if abogado else ""
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Matriz_Seguimiento{sufijo}_{hoy}.xlsx"},
    )


# ── Papelera de reciclaje (DEBE ir antes de "/{reg_id}" — ruta estática) ───────

@router.get("/papelera", response_class=HTMLResponse)
async def papelera(request: Request, msg: str = ""):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/matriz/?msg=sin_permiso", status_code=303)
    conn = get_db()
    where = ["eliminado_en IS NOT NULL"]
    params: list = []
    if user.get("rol") == "abogado":
        where.append("abogado = ?")
        params.append(user["nombre_completo"])
    rows = conn.execute(
        f"SELECT * FROM matriz_seguimiento WHERE {' AND '.join(where)} ORDER BY eliminado_en DESC",
        params,
    ).fetchall()
    conn.close()
    registros = [{
        "id": r["id"],
        "titulo": f"BPM {r['n_bpm']}" if r["n_bpm"] else (r["n_expediente"] or f"Trámite #{r['id']}"),
        "subtitulo": r["abogado"] or "",
        "eliminado_en": r["eliminado_en"],
        "eliminado_por": r["eliminado_por"],
    } for r in rows]
    return templates.TemplateResponse("papelera.html", tpl(request, _MOD,
        modulo_nombre="Matriz de Seguimiento", prefix="/matriz", base_template="base_matriz.html",
        registros=registros, volver_url="/matriz/", active="papelera", msg=msg,
    ))


@router.post("/{reg_id}/restaurar")
async def restaurar(request: Request, reg_id: int):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/matriz/papelera?msg=sin_permiso", status_code=303)
    conn = get_db()
    row = conn.execute("SELECT abogado FROM matriz_seguimiento WHERE id=?", (reg_id,)).fetchone()
    if not row or not _es_propietario(user, row["abogado"]):
        conn.close()
        return RedirectResponse("/matriz/papelera?msg=sin_permiso", status_code=303)
    conn.execute("UPDATE matriz_seguimiento SET eliminado_en = NULL, eliminado_por = NULL WHERE id = ?", (reg_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "restaurar", _MOD, f"Trámite #{reg_id}", registro_id=reg_id)
    return RedirectResponse("/matriz/papelera?msg=restaurado", status_code=303)


@router.post("/{reg_id}/purgar")
async def purgar(request: Request, reg_id: int):
    user = request.state.user
    if not user or user.get("rol") not in ROLES_SUPERUSUARIO:
        return RedirectResponse("/matriz/papelera?msg=sin_permiso", status_code=303)
    conn = get_db()
    conn.execute("DELETE FROM matriz_seguimiento WHERE id = ? AND eliminado_en IS NOT NULL", (reg_id,))
    conn.commit()
    conn.close()
    registrar_log(user, "purgar", _MOD, f"Trámite #{reg_id} — eliminado definitivamente", registro_id=reg_id)
    return RedirectResponse("/matriz/papelera?msg=purgado", status_code=303)


# ── Detalle ────────────────────────────────────────────────────────────────────

@router.get("/{reg_id}", response_class=HTMLResponse)
async def matriz_detalle(request: Request, reg_id: int, msg: str = ""):
    user = request.state.user
    conn = get_db()
    reg = conn.execute("SELECT * FROM matriz_seguimiento WHERE id = ?", (reg_id,)).fetchone()
    if not reg or not _es_propietario(user, reg["abogado"]):
        conn.close()
        return RedirectResponse("/matriz/?msg=no_encontrado")
    reg_dict = _enriquecer(dict(reg))
    historial = historial_registro(conn, _MOD, reg_id)
    conn.close()
    return templates.TemplateResponse("matriz_detalle.html", tpl(request, _MOD,
        reg=reg_dict, msg=msg, active="matriz_lista", historial=historial,
    ))


# ── Editar ─────────────────────────────────────────────────────────────────────

@router.get("/{reg_id}/editar", response_class=HTMLResponse)
async def matriz_editar_form(request: Request, reg_id: int):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse(f"/matriz/{reg_id}?msg=sin_permiso", status_code=303)
    conn = get_db()
    reg = conn.execute("SELECT * FROM matriz_seguimiento WHERE id = ?", (reg_id,)).fetchone()
    if not reg or not _es_propietario(user, reg["abogado"]):
        conn.close()
        return RedirectResponse("/matriz/?msg=no_encontrado")
    if reg["eliminado_en"]:
        conn.close()
        return RedirectResponse("/matriz/papelera?msg=error_en_papelera", status_code=303)
    abogados = _abogados_lista(conn)
    conn.close()
    es_abogado = user.get("rol") == "abogado"
    return templates.TemplateResponse("matriz_form.html", tpl(request, _MOD,
        reg=dict(reg), abogados=abogados, tipos=TIPOS_TRAMITE, etapas=ETAPAS_BPM, estados=ESTADOS,
        es_abogado=es_abogado, active="matriz_lista",
    ))


@router.post("/{reg_id}/editar")
async def matriz_editar_post(
    request: Request,
    reg_id: int,
    abogado: str = Form(""),
    n_expediente: str = Form(""),
    n_bpm: str = Form(""),
    tipo_tramite: str = Form(""),
    etapa_bpm: str = Form(""),
    asunto: str = Form(""),
    ultima_actuacion: str = Form(""),
    fecha_ultima_actuacion: str = Form(""),
    proxima_actuacion: str = Form(""),
    fecha_limite: str = Form(""),
    estado: str = Form(""),
    observaciones: str = Form(""),
):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse(f"/matriz/{reg_id}?msg=sin_permiso", status_code=303)

    conn = get_db()
    actual = conn.execute("SELECT abogado FROM matriz_seguimiento WHERE id=? AND eliminado_en IS NULL", (reg_id,)).fetchone()
    if not actual or not _es_propietario(user, actual["abogado"]):
        conn.close()
        return RedirectResponse(f"/matriz/{reg_id}?msg=sin_permiso", status_code=303)

    if user.get("rol") == "abogado":
        abogado = user["nombre_completo"]
    abogado = abogado.strip() or actual["abogado"]

    conn.execute(
        """UPDATE matriz_seguimiento
           SET abogado=?, n_expediente=?, n_bpm=?, tipo_tramite=?, etapa_bpm=?, asunto=?,
               ultima_actuacion=?, fecha_ultima_actuacion=?, proxima_actuacion=?, fecha_limite=?,
               estado=?, observaciones=?, updated_at=datetime('now','localtime')
           WHERE id=? AND eliminado_en IS NULL""",
        [
            abogado,
            n_expediente.strip() or None,
            n_bpm.strip() or None,
            tipo_tramite.strip() or None,
            etapa_bpm.strip() or None,
            asunto.strip() or None,
            ultima_actuacion.strip() or None,
            fecha_ultima_actuacion or None,
            proxima_actuacion.strip() or None,
            fecha_limite or None,
            estado.strip() or "EN TRÁMITE",
            observaciones.strip() or None,
            reg_id,
        ],
    )
    conn.commit()
    conn.close()
    registrar_log(user, "editar", _MOD, f"{abogado} — BPM {n_bpm or 's/n'} — {asunto or n_expediente or ''}", registro_id=reg_id)
    return RedirectResponse(f"/matriz/{reg_id}?msg=actualizado", status_code=303)


# ── Eliminar ───────────────────────────────────────────────────────────────────

@router.post("/{reg_id}/eliminar")
async def matriz_eliminar(request: Request, reg_id: int):
    user = request.state.user
    if not _pw(user, _MOD):
        return RedirectResponse("/matriz/?msg=sin_permiso", status_code=303)
    conn = get_db()
    actual = conn.execute("SELECT abogado FROM matriz_seguimiento WHERE id=?", (reg_id,)).fetchone()
    if not actual or not _es_propietario(user, actual["abogado"]):
        conn.close()
        return RedirectResponse("/matriz/?msg=sin_permiso", status_code=303)
    conn.execute(
        "UPDATE matriz_seguimiento SET eliminado_en = datetime('now','localtime'), eliminado_por = ? WHERE id = ?",
        (user.get("nombre_completo"), reg_id),
    )
    conn.commit()
    conn.close()
    registrar_log(user, "eliminar", _MOD, f"Trámite #{reg_id}",
                  request.client.host if request.client else None, registro_id=reg_id)
    return RedirectResponse("/matriz/?msg=eliminado", status_code=303)
